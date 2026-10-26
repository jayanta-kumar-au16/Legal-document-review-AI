"""Builds tiny but valid PDF and DOCX files for tests, with no extra dependencies beyond python-docx."""
import io


def make_pdf(pages: list[str]) -> bytes:
    """A real, text-based PDF: one Helvetica text block per page. Each string is a page; '\\n' starts a new line."""
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }
    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(f'{i} 0 R' for i in page_ids)}] /Count {len(pages)} >>".encode()

    for index, text in enumerate(pages):
        content_id = page_ids[index] + 1
        stream = b"BT /F1 11 Tf 14 TL 72 720 Td\n"
        for line in text.split("\n"):
            escaped = line.encode("cp1252", errors="replace").replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
            stream += b"(" + escaped + b") Tj T*\n"
        stream += b"ET"
        objects[page_ids[index]] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>"
        ).encode()
        objects[content_id] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += b"%d 0 obj\n" % number + objects[number] + b"\nendobj\n"
    xref_at = len(out)
    size = max(objects) + 1
    out += b"xref\n0 %d\n0000000000 65535 f \n" % size
    for number in range(1, size):
        out += b"%010d 00000 n \n" % offsets[number]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref_at)
    return bytes(out)


def make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# The fictional sample NDA from the spec, split across two pages like the seeded one.
SAMPLE_PAGE_1 = """NON-DISCLOSURE AGREEMENT
Fictional sample for software demonstration only
This Agreement is entered into between Company A ("Disclosing Party") and Company B ("Receiving Party").
Clause 1 — Purpose
The Receiving Party may use confidential information solely to evaluate a potential business relationship with the Disclosing Party.
Clause 2 — Confidentiality
The Receiving Party shall protect all confidential information and shall not disclose it to third parties without prior written consent.
Clause 3 — Duration
The obligations under this agreement shall continue for one year from the date of disclosure."""

SAMPLE_PAGE_2 = """Clause 4 — Return of Information
The Receiving Party shall return confidential information upon written request from the Disclosing Party.
Clause 5 — Governing Law
This agreement shall be governed by applicable law."""

EXPECTED_SAMPLE_OUTCOMES = {
    "R1": "pass", "R2": "pass", "R3": "review_required", "R4": "deviation", "R5": "deviation",
}

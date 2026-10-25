import io
from uuid import uuid4
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from reportlab.pdfgen import canvas
from app.main import create_app, DOCX_MIME


def pdf_bytes(text='Confidential information shall be protected.', blank_second=False):
    out = io.BytesIO()
    c = canvas.Canvas(out)
    c.drawString(50, 750, text)
    c.showPage()
    if blank_second:
        c.showPage()
    c.save()
    return out.getvalue()

@pytest.fixture
def client(tmp_path):
    return TestClient(create_app(tmp_path))

def test_pdf_roundtrip_and_restart(client, tmp_path):
    raw = pdf_bytes()
    response = client.post('/api/documents', files={'file': ('nda.pdf', raw, 'application/pdf')})
    assert response.status_code == 201
    identifier = response.json()['document_id']
    text = TestClient(create_app(tmp_path)).get(f'/api/documents/{identifier}/text').json()
    assert text['segments'][0]['page'] == 1
    assert 'Confidential information' in text['segments'][0]['text']
    assert (tmp_path / identifier / 'original.pdf').read_bytes() == raw

def test_docx_order_and_references(client):
    d = Document()
    d.add_paragraph('Before table')
    d.add_table(rows=1, cols=1).cell(0,0).text = 'Table evidence'
    d.add_paragraph('After table')
    out = io.BytesIO(); d.save(out)
    r = client.post('/api/documents', files={'file': ('nda.docx', out.getvalue(), DOCX_MIME)})
    assert r.status_code == 201
    result = client.get(f"/api/documents/{r.json()['document_id']}/text").json()
    assert [s['text'] for s in result['segments']] == ['Before table','Table evidence','After table']
    assert result['page_count'] is None
    assert result['segments'][1]['table'] == 1
    assert all(s['page'] is None for s in result['segments'])

@pytest.mark.parametrize('name,body,mime,status', [
 ('nda.txt', b'hello', 'text/plain',415),
 ('nda.pdf', b'', 'application/pdf',422),
 ('nda.pdf', b'not PDF', 'application/pdf',422),
 ('nda.docx', b'not DOCX', DOCX_MIME,422),
 ('nda.pdf', b'hello', 'image/png',415),
])
def test_invalid_files(client,name,body,mime,status):
    assert client.post('/api/documents',files={'file':(name,body,mime)}).status_code == status

def test_size_limit(tmp_path):
    client = TestClient(create_app(tmp_path, max_upload_bytes=10))
    assert client.post('/api/documents',files={'file':('nda.pdf',b'x'*11,'application/pdf')}).status_code == 413
    assert list(tmp_path.iterdir()) == []

def test_blank_pdf(client):
    w=PdfWriter(); w.add_blank_page(width=100,height=100)
    out=io.BytesIO(); w.write(out)
    r=client.post('/api/documents',files={'file':('blank.pdf',out.getvalue(),'application/pdf')})
    assert r.status_code == 422
    assert r.json()['detail']['code'] == 'no_extractable_text'

def test_partial_pdf_warning(client):
    r=client.post('/api/documents',files={'file':('partial.pdf',pdf_bytes(blank_second=True),'application/pdf')})
    assert r.status_code == 201
    assert r.json()['page_count'] == 2
    assert r.json()['extracted_page_count'] == 1
    assert r.json()['extraction_status'] == 'completed_with_warnings'

def test_encrypted_pdf(client):
    w=PdfWriter(); w.add_blank_page(width=100,height=100); w.encrypt('secret')
    out=io.BytesIO(); w.write(out)
    r=client.post('/api/documents',files={'file':('encrypted.pdf',out.getvalue(),'application/pdf')})
    assert r.status_code == 422
    assert r.json()['detail']['code'] == 'encrypted_pdf'

def test_unknown_and_invalid_id(client):
    assert client.get(f'/api/documents/{uuid4()}/text').status_code == 404
    assert client.get('/api/documents/not-a-uuid/text').status_code == 422

def test_missing_file(client):
    assert client.post('/api/documents').status_code == 422

def test_sample_nda(client):
    from pathlib import Path
    raw = Path('test-data/sample_nda.pdf').read_bytes()
    r = client.post('/api/documents', files={'file':('sample_nda.pdf',raw,'application/pdf')})
    assert r.status_code == 201
    result = client.get(f"/api/documents/{r.json()['document_id']}/text").json()
    text = '\n'.join(s['text'] for s in result['segments'])
    for title in ['Purpose','Confidentiality','Duration','Return of Information','Governing Law']:
        assert title in text
    assert 'one year from the date of disclosure' in text
